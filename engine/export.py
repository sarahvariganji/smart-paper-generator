"""Stage 4 — export the paper as a Word document."""
import io, math
from .bank import BY_ID, MK, SECTIONS

def to_docx(ids, answers=False):
    from docx import Document  # lazy import: the engine itself needs no third-party packages
    paper = [BY_ID[i] for i in ids]
    total = sum(MK[q["y"]] for q in paper)
    d = Document()
    d.add_heading("Mathematics — Question Paper", 0)
    d.add_paragraph(f"Time: {math.ceil(total * 1.5)} minutes        Maximum marks: {total}")
    d.add_paragraph("Attempt all questions. Show your working for Sections B and C.")
    n = 0
    for y, title in SECTIONS:
        qs = [q for q in paper if q["y"] == y]
        if not qs:
            continue
        d.add_heading(title, 2)
        for q in qs:
            n += 1
            d.add_paragraph(f"{n}. {q['q']}  [{MK[y]}]")
            if q["o"]:
                d.add_paragraph("      ".join(f"({'abcd'[j]}) {o}" for j, o in enumerate(q["o"])))
            if answers:
                a = f"({'abcd'[q['a']]}) {q['o'][q['a']]}" if q["o"] else q["a"]
                d.add_paragraph(f"Answer: {a}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
